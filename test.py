import pandas as pd
from docx import Document
import os
import re

# --- MAPPING CONFIGURATION ---
# Update these keys to match your EXACT Word filenames in the folder
FILE_TO_BOOK_MAP = {
    "Book_One_Part_One_of_the_Book_of_Negroes.docx": "Book One Part One",
    "Book_One_Part_Two_of_the_Book_of_Negroes.docx": "Book One Part Two",
    "Book_Two.docx": "Book Two",
    "Book_Three.docx": "Book Three"
}

def extract_ship_name(text):
    """
    Extracts ship name based on boarding/destination keywords.
    Ensures 'Spring bound' becomes 'Spring' by stopping at keywords.
    """
    # 1. Patterns that look for a name FOLLOWED by a keyword (bound, headed, etc.)
    # The ([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*?) is non-greedy to stop exactly at the keyword
    keyword_patterns = [
        r"(?:Ship|Brig|Sloop|the)\s+([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*?)\s+(?:bound|headed|destination|for|master|commander|capt)",
        r"([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*?)\s+(?:bound|headed|for)\s+to",
    ]
    
    # 2. Standard patterns if no 'bound/for' keywords are present
    standard_patterns = [
        r"in the\s+([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*)", 
        r"board the\s+([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*)",
        r"Ship\s+([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*)",
        r"Brig\s+([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*)",
        r"Sloop\s+([A-Z][a-z]+(?:\s+[A-Z][a-z]+)*)"
    ]
    
    # Try keywords first to ensure "Spring bound" -> "Spring"
    for pattern in keyword_patterns:
        match = re.search(pattern, text, re.I)
        if match:
            return match.group(1).strip()
            
    # Fallback to standard patterns
    for pattern in standard_patterns:
        match = re.search(pattern, text, re.I)
        if match:
            return match.group(1).strip()
            
    return "Unknown/Not Found"

def get_word_content(directory_path):
    """Scrapes text and tracks filenames from the Word directory."""
    raw_data = []
    if not os.path.exists(directory_path):
        print(f"Error: Folder '{directory_path}' not found.")
        return []

    for filename in os.listdir(directory_path):
        if filename.endswith(".docx") and not filename.startswith("~$"):
            file_path = os.path.join(directory_path, filename)
            try:
                doc = Document(file_path)
                entries = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                entries.append(cell.text.strip())
                
                for entry in entries:
                    raw_data.append({"Notes": entry, "Source_Word_File": filename})
            except Exception as e:
                print(f"Error reading {filename}: {e}")
    return raw_data

def process_loyallist_comparison(excel_path, word_dir, output_file):
    print(f"Loading {excel_path}...")
    try:
        # Load the master consolidated file
        df_master = pd.read_excel(excel_path)
    except Exception as e:
        print(f"Excel Load Error: {e}")
        return

    # Clean the Master Excel data for comparison
    df_master['Ship_Name'] = df_master['Ship_Name'].astype(str).str.strip().str.lower()
    df_master['Book'] = df_master['Book'].astype(str).str.strip()
    df_master['Notes'] = df_master['Notes'].fillna("").astype(str).str.strip()

    # Build Lookup set for existing Ship+Book combinations
    df_master['Ship_Book_Key'] = df_master['Ship_Name'] + "||" + df_master['Book']
    existing_ship_book_combos = set(df_master['Ship_Book_Key'].unique())
    
    # Build Lookup set for Notes
    existing_notes = set(df_master['Notes'].unique())

    print(f"Scraping Word documents in '{word_dir}'...")
    word_records = get_word_content(word_dir)

    missing_records = []
    seen_in_report = set()

    for record in word_records:
        note_text = record["Notes"]
        source_file = record["Source_Word_File"]
        mapped_book = FILE_TO_BOOK_MAP.get(source_file, "Unknown Book")
        
        # 1. Extract Ship Name (Logic improved to stop at 'bound')
        extracted_ship = extract_ship_name(note_text)
        
        # 2. VALIDATION STEP 1: Check Ship Name First
        # Normalize to lower for comparison
        ship_book_key = f"{extracted_ship.lower()}||{mapped_book}"
        
        if extracted_ship != "Unknown/Not Found" and ship_book_key in existing_ship_book_combos:
            # Skip this record because the ship is already recorded in this Book
            continue
            
        # 3. VALIDATION STEP 2: Check if the exact Note text exists
        if note_text in existing_notes:
            continue

        # 4. If both checks fail, it is truly missing
        report_key = (note_text, source_file, extracted_ship)
        if report_key not in seen_in_report:
            missing_records.append({
                "Notes": note_text,
                "Source_Word_File": source_file,
                "Ship": extracted_ship
            })
            seen_in_report.add(report_key)

    if missing_records:
        df_final = pd.DataFrame(missing_records)
        df_final = df_final[["Notes", "Source_Word_File", "Ship"]]
        df_final.to_excel(output_file, index=False)
        print(f"\nAnalysis complete. Found {len(missing_records)} records missing from the Master Excel.")
        print(f"Report saved as: {output_file}")
    else:
        print("\nAll entries accounted for. No missing records found based on Ship/Book validation.")

# --- File Paths ---
DIR_NAME = "Book_of_Negroes"
MASTER_FILE = "Black_Loyalist_Directory_Consolidated.xlsx"
REPORT_FILE = "Validated_Missing_Records.xlsx"

if __name__ == "__main__":
    process_loyallist_comparison(MASTER_FILE, DIR_NAME, REPORT_FILE)